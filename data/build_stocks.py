"""Turn the lab stock list into a file the app can load.

    python3 build_stocks.py

Reads ../../JRDrosophilaStocksJune2026.xlsx and writes my-stocks.js, the file the
app loads at startup if it is there. data/stocks.js is not touched: that holds
the standard balancer stocks every copy ships with. The app runs straight off the
disk, where fetch() is blocked, so the data has to arrive as a script rather than
as JSON.
"""
import json
import os
import re

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
BOOK = os.path.join(HERE, '..', '..', 'JRDrosophilaStocksJune2026.xlsx')
OUT = os.path.join(HERE, 'my-stocks.js')


def clean(v, limit=160):
    if v is None:
        return ''
    t = re.sub(r'\s+', ' ', str(v)).strip()
    return t[:limit]


def main():
    wb = openpyxl.load_workbook(BOOK, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        number, name, desc, project, more, in_stock = (list(r) + [None] * 6)[:6]
        genotype = clean(name)
        if not genotype:
            continue
        rows.append({
            'n': clean(number, 24),
            'g': genotype,
            'd': clean(desc, 120),
            'p': clean(project, 40),
            's': 1 if clean(in_stock).lower() == 'yes' else 0,
        })

    held = sum(r['s'] for r in rows)
    body = ',\n'.join(json.dumps(r, ensure_ascii=False, separators=(',', ':')) for r in rows)
    name = os.path.basename(BOOK)
    with open(OUT, 'w') as fh:
        fh.write('/* A stock list for the Fly Cross Planner, built from %s\n' % name)
        fh.write(' * by build_stocks.py: %d rows, %d of them in stock. This file is personal to\n' % (len(rows), held))
        fh.write(' * one lab - a clean copy of the app ships it empty. Do not edit by hand; edit\n')
        fh.write(' * the spreadsheet and run the script again.\n')
        fh.write(' *   n = stock number, g = genotype, d = description, p = project, s = in stock\n')
        fh.write(' */\n')
        fh.write("(function (root) {\n  'use strict';\n  var FCS = root.FCS = root.FCS || {};\n")
        fh.write('  FCS.myStocks = { name: %s, rows: [\n' % json.dumps(name))
        fh.write(body)
        fh.write('\n  ] };\n})(typeof globalThis !== \'undefined\' ? globalThis : this);\n')
    print('wrote %d stocks (%d held) to %s' % (len(rows), held, OUT))


if __name__ == '__main__':
    main()
